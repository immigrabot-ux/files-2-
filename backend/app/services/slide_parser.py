from pptx import Presentation
from PyPDF2 import PdfReader
from PIL import Image
from io import BytesIO
from loguru import logger
from typing import List, Dict, Optional
import tempfile
import os


class Slide:
    """Represents a single slide"""
    def __init__(self, number: int, text: str, image_data: Optional[bytes] = None, speaker_notes: Optional[str] = None):
        self.number = number
        self.text = text
        self.image_data = image_data
        self.speaker_notes = speaker_notes


def parse_powerpoint(file_data: bytes) -> List[Slide]:
    """
    Parse PowerPoint file and extract slides.

    Args:
        file_data: PowerPoint file content as bytes

    Returns:
        List of Slide objects
    """
    try:
        # Save to temp file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.pptx') as tmp:
            tmp.write(file_data)
            tmp_path = tmp.name

        prs = Presentation(tmp_path)
        slides = []

        for idx, slide in enumerate(prs.slides, start=1):
            # Extract text from shapes
            text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text_parts.append(shape.text)

            slide_text = "\n".join(text_parts)

            # Extract speaker notes
            speaker_notes = None
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                notes_text_frame = notes_slide.notes_text_frame
                if notes_text_frame and notes_text_frame.text:
                    speaker_notes = notes_text_frame.text

            # Convert slide to image
            # Note: python-pptx doesn't directly export to image
            # In production, use a library like pdf2image or LibreOffice headless
            image_data = None  # Placeholder

            slides.append(Slide(
                number=idx,
                text=slide_text,
                image_data=image_data,
                speaker_notes=speaker_notes
            ))

        # Cleanup
        os.unlink(tmp_path)

        logger.info(f"Parsed PowerPoint: {len(slides)} slides")
        return slides

    except Exception as exc:
        logger.error(f"PowerPoint parsing error: {str(exc)}")
        raise


def parse_pdf(file_data: bytes) -> List[Slide]:
    """
    Parse PDF file and extract slides.

    Args:
        file_data: PDF file content as bytes

    Returns:
        List of Slide objects
    """
    try:
        # Save to temp file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
            tmp.write(file_data)
            tmp_path = tmp.name

        reader = PdfReader(tmp_path)
        slides = []

        for idx, page in enumerate(reader.pages, start=1):
            # Extract text
            text = page.extract_text()

            # Convert page to image
            # In production, use pdf2image library
            image_data = None  # Placeholder

            slides.append(Slide(
                number=idx,
                text=text,
                image_data=image_data,
                speaker_notes=None
            ))

        # Cleanup
        os.unlink(tmp_path)

        logger.info(f"Parsed PDF: {len(slides)} slides")
        return slides

    except Exception as exc:
        logger.error(f"PDF parsing error: {str(exc)}")
        raise


def parse_slides(file_data: bytes, file_type: str) -> List[Dict]:
    """
    Parse slides from file.

    Args:
        file_data: File content as bytes
        file_type: File extension (.pptx, .pdf, .key)

    Returns:
        List of slide dictionaries
    """
    try:
        if file_type == '.pptx':
            slides = parse_powerpoint(file_data)
        elif file_type == '.pdf':
            slides = parse_pdf(file_data)
        elif file_type == '.key':
            # Keynote files need special handling (convert to PDF first)
            raise NotImplementedError("Keynote support coming soon")
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

        # Convert to dict format
        slide_dicts = []
        for slide in slides:
            slide_dicts.append({
                "number": slide.number,
                "text": slide.text,
                "speakerNotes": slide.speaker_notes,
                "imageData": slide.image_data
            })

        return slide_dicts

    except Exception as exc:
        logger.error(f"Slide parsing error: {str(exc)}")
        raise


def validate_file(file_data: bytes, file_type: str, max_size_mb: int = 100) -> tuple[bool, Optional[str]]:
    """
    Validate uploaded file.

    Args:
        file_data: File content as bytes
        file_type: File extension
        max_size_mb: Maximum file size in MB

    Returns:
        Tuple of (is_valid, error_message)
    """
    # Check file type
    if file_type not in ['.pptx', '.pdf', '.key']:
        return False, "Only .pptx, .pdf, and .key files are supported"

    # Check file size
    size_mb = len(file_data) / (1024 * 1024)
    if size_mb > max_size_mb:
        return False, f"File size exceeds {max_size_mb}MB limit"

    # Check if file is empty
    if len(file_data) == 0:
        return False, "File is empty"

    return True, None
